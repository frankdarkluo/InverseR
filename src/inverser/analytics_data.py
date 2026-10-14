import logging
import math
import re
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Literal

from pydantic import Field

from inverser.runtime import file_sha256, fingerprint, write_json
from inverser.schemas import Record

LOGGER = logging.getLogger(__name__)
SPREADSHEET_NS = "urn:schemas-microsoft-com:office:spreadsheet"
NS = {"s": SPREADSHEET_NS}


class AnnotationPoint(Record):
    source_marker: str
    xy_source: tuple[float, float]
    time_text: str
    annotation_video_time_s: float | None = Field(default=None, ge=0)


class AnalyticsLayout(Record):
    schema_version: Literal["analytics-layout-v1"] = "analytics-layout-v1"
    sample_id: str
    role: Literal["auxiliary_real_layout"] = "auxiliary_real_layout"
    game: Literal["nine_ball"] = "nine_ball"
    source_file: str
    source_sha256: str
    source_group: str
    source_tournament: str
    source_match: str
    source_rack_label: str
    video_title: str | None
    points: tuple[AnnotationPoint, ...] = Field(min_length=1)
    coordinate_basis: Literal["source_native_units_unverified"] = "source_native_units_unverified"
    marker_to_ball_identity_verified: Literal[False] = False
    complete_inventory_verified: Literal[False] = False
    measured_cue_action_available: Literal[False] = False
    measured_spin_available: Literal[False] = False
    publication_annotation_not_new_human_review: Literal[True] = True
    main_snooker_benchmark_eligible: Literal[False] = False
    flags: tuple[str, ...]


class TrajectorySample(Record):
    xy_cm: tuple[float, float]
    time_text: str
    annotation_video_time_s: float = Field(ge=0)


class AnalyticsTrajectory(Record):
    schema_version: Literal["analytics-trajectory-v1"] = "analytics-trajectory-v1"
    sample_id: str
    role: Literal["auxiliary_real_factual_trajectory"] = "auxiliary_real_factual_trajectory"
    game: Literal["nine_ball"] = "nine_ball"
    source_file: str
    source_sha256: str
    source_group: str
    source_label: str
    source_video_title: str
    coordinate_basis: Literal["source_declared_centimetres"] = "source_declared_centimetres"
    source_spatial_calibration_independently_verified: Literal[False] = False
    samples: tuple[TrajectorySample, ...] = Field(min_length=2)
    timestamps_strictly_increasing: bool
    complete_scene_inventory_available: Literal[False] = False
    measured_cue_action_available: Literal[False] = False
    measured_spin_available: Literal[False] = False
    action_conditioned_intervention_ground_truth: Literal[False] = False
    main_snooker_benchmark_eligible: Literal[False] = False


def spreadsheet_xml_rows(path: Path) -> tuple[tuple[str, tuple[tuple[str, ...], ...]], ...]:
    if path.stat().st_size > 10_000_000:
        raise ValueError("Spreadsheet XML exceeds the bounded annotation-file limit")
    payload = path.read_bytes()
    if b"<!DOCTYPE" in payload.upper() or b"<!ENTITY" in payload.upper():
        raise ValueError("Spreadsheet annotation must not contain entity declarations")
    root = ET.fromstring(payload)
    if root.tag != f"{{{SPREADSHEET_NS}}}Workbook":
        raise ValueError("Annotation is not the expected source SpreadsheetML workbook")
    sheets = []
    for sheet in root.findall("s:Worksheet", NS):
        rows = []
        for row in sheet.findall("s:Table/s:Row", NS):
            values: list[str] = []
            for cell in row.findall("s:Cell", NS):
                index = cell.get(f"{{{SPREADSHEET_NS}}}Index")
                if index is not None:
                    position = int(index) - 1
                    if not len(values) <= position <= 1000:
                        raise ValueError("Invalid source spreadsheet sparse-column index")
                    values.extend([""] * (position - len(values)))
                values.append(cell.findtext("s:Data", default="", namespaces=NS))
            rows.append(tuple(values))
        sheets.append((sheet.get(f"{{{SPREADSHEET_NS}}}Name", ""), tuple(rows)))
    if not sheets:
        raise ValueError("Spreadsheet annotation has no worksheets")
    return tuple(sheets)


def annotation_time_seconds(text: str, *, allow_seconds_only: bool = False) -> float | None:
    if not text:
        return None
    if text.startswith("1899-12-31T"):
        value = datetime.fromisoformat(text)
        if value.tzinfo is not None:
            raise ValueError("Source spreadsheet time zone is not verified")
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
    if allow_seconds_only and re.fullmatch(r"\d{1,2}\.\d+", text):
        seconds = float(text)
        if seconds >= 60:
            raise ValueError("Seconds-only source clock exceeds its minute range")
        return seconds
    if not re.fullmatch(r"\d+:\d{1,2}(?:\.\d+)?|\d+:\d{1,2}:\d{1,2}(?:\.\d+)?", text):
        raise ValueError("Unsupported source annotation clock format")
    parts = text.split(":")
    seconds = float(parts[-1])
    if not seconds < 60:
        raise ValueError("Source annotation seconds lie outside the clock range")
    minutes = int(parts[-2])
    if len(parts) == 3 and minutes >= 60:
        raise ValueError("Source annotation minutes lie outside the clock range")
    return seconds + minutes * 60 + (int(parts[0]) * 3600 if len(parts) == 3 else 0)


def read_analytics_layout(path: Path, coordinate_root: Path) -> AnalyticsLayout:
    relative = path.resolve().relative_to(coordinate_root.resolve())
    if len(relative.parts) < 3:
        raise ValueError("Layout provenance requires tournament/match/rack paths")
    points = []
    flags = []
    title = None
    for name, rows in spreadsheet_xml_rows(path):
        title = title or name
        in_points = False
        for row in rows:
            if tuple(value.strip() for value in row[:4]) == ("Name", "X", "Y", "Time"):
                in_points = True
                continue
            if not in_points:
                continue
            if not row or not any(row):
                in_points = False
                continue
            if len(row) < 4:
                raise ValueError("Layout point row lacks source coordinates or clock")
            try:
                xy = float(row[1]), float(row[2])
            except ValueError as error:
                raise ValueError("Layout annotation has a nonnumeric coordinate") from error
            if not all(math.isfinite(value) for value in xy):
                raise ValueError("Layout annotation has a non-finite coordinate")
            try:
                time = annotation_time_seconds(row[3])
            except ValueError:
                LOGGER.debug("Unverified annotation clock in %s: %r", path.name, row[3])
                time = None
                flags.append("unparsed_source_clock")
            points.append(
                AnnotationPoint(
                    source_marker=row[0],
                    xy_source=xy,
                    time_text=row[3],
                    annotation_video_time_s=time,
                )
            )
    if not points:
        raise ValueError("Source layout has no declared Name/X/Y/Time point table")
    if len(points) != 10:
        flags.append("non_ten_marker_layout")
    if len({point.source_marker for point in points}) != len(points):
        flags.append("repeated_source_marker_names")
    if any(min(point.xy_source) < 0 for point in points):
        flags.append("negative_source_coordinate_status_unverified")
    if len({point.time_text for point in points}) != 1:
        flags.append("multiple_annotation_times_not_one_snapshot")
    years = re.findall(r"\b(?:19|20)\d{2}\b", relative.parts[0])
    if title and years and years[0] not in title:
        flags.append("directory_year_not_in_source_video_title")
    checksum = file_sha256(path)
    return AnalyticsLayout(
        sample_id="bsa-layout-" + fingerprint((str(relative), checksum))[:20],
        source_file=str(path.resolve()),
        source_sha256=checksum,
        source_group=fingerprint(relative.parts[:2]),
        source_tournament=relative.parts[0],
        source_match=relative.parts[1],
        source_rack_label=path.stem,
        video_title=title,
        points=tuple(points),
        flags=tuple(sorted(set(flags))),
    )


def read_analytics_trajectory(path: Path, root: Path) -> AnalyticsTrajectory:
    sheets = spreadsheet_xml_rows(path)
    if len(sheets) != 1:
        raise ValueError("Factual trajectory requires one source-native track worksheet")
    title, rows = sheets[0]
    label = None
    centimetres_declared = False
    start = None
    for index, row in enumerate(rows):
        if len(row) >= 2 and row[0].strip() == "Label :":
            label = row[1].strip()
        if any("x,y:cm" in value.replace(" ", "").lower() for value in row):
            centimetres_declared = True
        if tuple(value.strip().lower() for value in row[:3]) == ("x", "y", "t"):
            start = index + 1
            break
    if not label or start is None or not centimetres_declared:
        raise ValueError("Trajectory lacks explicit source label, centimetre units or x/y/t table")
    samples = []
    for row in rows[start:]:
        if not row or not any(row):
            break
        if len(row) < 3:
            raise ValueError("Trajectory point row lacks coordinates or source time")
        xy = float(row[0]), float(row[1])
        if not all(math.isfinite(value) for value in xy):
            raise ValueError("Trajectory contains non-finite source coordinates")
        time = annotation_time_seconds(row[2], allow_seconds_only=True)
        if time is None:
            raise ValueError("Factual timed trajectory cannot have missing source timestamps")
        samples.append(TrajectorySample(xy_cm=xy, time_text=row[2], annotation_video_time_s=time))
    relative = path.resolve().relative_to(root.resolve())
    checksum = file_sha256(path)
    return AnalyticsTrajectory(
        sample_id="bsa-trajectory-" + fingerprint((str(relative), checksum))[:20],
        source_file=str(path.resolve()),
        source_sha256=checksum,
        source_group=fingerprint(relative.parts[:-1]),
        source_label=label,
        source_video_title=title,
        samples=tuple(samples),
        timestamps_strictly_increasing=all(
            a.annotation_video_time_s < b.annotation_video_time_s
            for a, b in zip(samples, samples[1:])
        ),
    )


def audit_analytics_data(
    dataset: Path, output: Path, trajectories: Path | None = None
) -> dict[str, object]:
    from openpyxl import load_workbook

    coordinates = dataset / "Coordinates"
    variables = dataset / "Variables"
    if not coordinates.is_dir() or not variables.is_dir():
        raise ValueError("Analytics release lacks its declared coordinate/variable folders")
    output.mkdir(parents=True, exist_ok=False)
    layouts = []
    errors = []
    for path in sorted(coordinates.rglob("*.xml")):
        try:
            layout = read_analytics_layout(path, coordinates)
        except (ValueError, ET.ParseError) as error:
            LOGGER.error("Layout annotation %s rejected: %s", path, error)
            errors.append({"file": str(path), "error": str(error)})
        else:
            layouts.append(layout)
    with (output / "layouts.jsonl").open("w", encoding="utf-8") as stream:
        for layout in layouts:
            stream.write(layout.model_dump_json() + "\n")
    workbook_rows = []
    office_lock_files = [
        path for path in sorted(variables.rglob("*.xlsx")) if path.name.startswith("~$")
    ]
    variable_files = [
        path for path in sorted(variables.rglob("*.xlsx")) if not path.name.startswith("~$")
    ]
    for path in variable_files:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook:
                rows = tuple(
                    tuple("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(values_only=True)
                )
                headers = [
                    row
                    for row in rows[:10]
                    if any(
                        value.strip().lower() in ("frame", "frame ", "who break") for value in row
                    )
                ]
                workbook_rows.append(
                    {
                        "source_file": str(path),
                        "source_sha256": file_sha256(path),
                        "sheet": sheet.title,
                        "rows": len(rows),
                        "headers": headers,
                        "formulas_evaluated": False,
                    }
                )
        finally:
            workbook.close()
    trace_rows = []
    typed_trajectories = []
    if trajectories is not None:
        for path in sorted(trajectories.rglob("*.xml")):
            try:
                sheets = spreadsheet_xml_rows(path)
            except (ValueError, ET.ParseError) as error:
                LOGGER.error("Trajectory annotation %s rejected: %s", path, error)
                errors.append({"file": str(path), "error": str(error)})
                continue
            trace_rows.append(
                {
                    "source_file": str(path),
                    "source_sha256": file_sha256(path),
                    "sheet_summaries": [
                        {"name": name, "rows": len(rows), "first_rows": rows[:12]}
                        for name, rows in sheets
                    ],
                    "actual_time_and_ball_scope_require_audit": True,
                    "measured_cue_controls_assumed": False,
                }
            )
            try:
                trajectory = read_analytics_trajectory(path, trajectories)
            except ValueError as error:
                LOGGER.error("Timed trajectory %s not eligible: %s", path, error)
                errors.append({"file": str(path), "error": str(error)})
            else:
                typed_trajectories.append(trajectory)
    with (output / "trajectories.jsonl").open("w", encoding="utf-8") as stream:
        for trajectory in typed_trajectories:
            stream.write(trajectory.model_dump_json() + "\n")
    write_json(output / "workbook_inventory.json", workbook_rows)
    write_json(output / "trajectory_table_inventory.json", trace_rows)
    write_json(output / "parse_errors.json", errors)
    summary: dict[str, object] = {
        "status": "complete" if not errors else "partial_with_explicit_annotation_errors",
        "coordinate_xml_files": len(list(coordinates.rglob("*.xml"))),
        "parsed_layout_records": len(layouts),
        "unique_layout_file_checksums": len({layout.source_sha256 for layout in layouts}),
        "source_match_groups": len({layout.source_group for layout in layouts}),
        "source_tournament_directories": len({layout.source_tournament for layout in layouts}),
        "marker_count_distribution": dict(Counter(len(layout.points) for layout in layouts)),
        "layout_flag_counts": dict(Counter(flag for layout in layouts for flag in layout.flags)),
        "variable_workbooks": len(variable_files),
        "ignored_office_lock_files": len(office_lock_files),
        "trajectory_xml_tables_acquired": len(trace_rows),
        "parsed_factual_trajectories": len(typed_trajectories),
        "trajectory_source_groups": len({track.source_group for track in typed_trajectories}),
        "trajectory_source_label_counts": dict(
            Counter(track.source_label for track in typed_trajectories)
        ),
        "strict_timed_trajectories": sum(
            track.timestamps_strictly_increasing for track in typed_trajectories
        ),
        "trajectory_source_samples": sum(len(track.samples) for track in typed_trajectories),
        "parse_errors": len(errors),
        "game": "nine_ball_auxiliary",
        "new_main_snooker_tasks": 0,
        "coordinate_units_independently_verified": False,
        "marker_ball_mapping_independently_verified": False,
        "measured_cue_speed_or_spin_records": 0,
        "downloaded_scripts_executed": False,
        "downloaded_weights_loaded": False,
        "annotation_timestamps_not_physical_action_measurements": True,
    }
    unresolved_clocks = sum("unparsed_source_clock" in layout.flags for layout in layouts)
    if unresolved_clocks:
        LOGGER.warning(
            "%d layout files have unverified source clocks; raw expressions and flags retained",
            unresolved_clocks,
        )
    write_json(output / "ignored_office_metadata.json", [str(path) for path in office_lock_files])
    write_json(output / "summary.json", summary)
    return summary
